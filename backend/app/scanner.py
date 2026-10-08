import hashlib
import io
import json
import os
import re
import stat
import zipfile
from pathlib import Path
from typing import TypedDict

from docx import Document
from pypdf import PdfReader

MAX_FILE = 20 * 1024 * 1024
MAX_BATCH = 100 * 1024 * 1024
MAX_FILES = 500
MAX_ENTRIES = 5000
EXCLUDED = {
    ".git",
    ".venv",
    "node_modules",
    ".next",
    "dist",
    "build",
    "data",
    "__pycache__",
    ".pytest_cache",
    ".mypy_cache",
    ".ruff_cache",
    "coverage",
}
DEPENDENCIES = {"package.json", "pyproject.toml", "requirements.txt", "Cargo.toml"}
FORMATS = {".md", ".txt", ".docx", ".pdf"}
SECRET = re.compile(
    r"-----BEGIN [A-Z ]*PRIVATE KEY-----|\bsk-[A-Za-z0-9_-]{16,}|"
    r"(?i:api[_-]?key|access[_-]?key|password|secret)\s*[:=]\s*['\"]?[^\s'\"]{8,}"
)


class FileRecord(TypedDict):
    path: str
    size: int
    sha256: str
    status: str
    chars: int


class BusinessError(Exception):
    def __init__(self, code: str, message: str, status: int = 400):
        self.code, self.message, self.status = code, message, status


def digest(value) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()
    ).hexdigest()


def canonical(raw: str, roots: tuple[Path, ...]) -> Path:
    path = Path(raw)
    if not path.is_absolute() or ".." in path.parts:
        raise BusinessError("PATH_FORBIDDEN", "请填写授权根目录内的绝对路径")
    try:
        resolved = path.resolve()
    except (OSError, RuntimeError):
        raise BusinessError("PATH_FORBIDDEN", "无法安全解析目录") from None
    if not any(resolved.is_relative_to(root) for root in roots):
        raise BusinessError("PATH_FORBIDDEN", "目录超出已授权范围")
    if any(part in EXCLUDED - {"data"} or part.startswith(".env") for part in resolved.parts):
        raise BusinessError("PATH_FORBIDDEN", "该目录属于受保护的秘密或依赖区域")
    return resolved


def read_safe(path: Path, dir_fd: int | None = None) -> bytes:
    # O_NOFOLLOW refuses last-component symlinks; callers recheck canonical roots.
    fd = os.open(path.name if dir_fd is not None else path, os.O_RDONLY | os.O_NOFOLLOW, dir_fd=dir_fd)
    with os.fdopen(fd, "rb") as file:
        metadata = os.fstat(file.fileno())
        if not stat.S_ISREG(metadata.st_mode):
            raise ValueError("not regular")
        if metadata.st_size > MAX_FILE:
            raise BusinessError("SCAN_LIMIT", "单文档超过 20 MB，请分批或缩小资料范围")
        raw = file.read(MAX_FILE + 1)
    if len(raw) > MAX_FILE:
        raise BusinessError("SCAN_LIMIT", "单文档超过 20 MB")
    return raw


def parse(raw: bytes, suffix: str) -> tuple[str, int]:
    try:
        if suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_BATCH:
                    return "unsafe_archive", 0
                if len(archive.infolist()) > 2000:
                    return "unsafe_archive", 0
                names = archive.namelist()
                if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                    return "invalid_type", 0
                if any(".." in Path(n).parts or n.startswith("/") for n in names):
                    return "unsafe_archive", 0
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        elif suffix == ".pdf":
            if not raw.startswith(b"%PDF-"):
                return "invalid_type", 0
            reader = PdfReader(io.BytesIO(raw), strict=True)
            if reader.is_encrypted or len(reader.pages) > 500:
                return "needs_manual", 0
            text = "\n".join(p.extract_text() or "" for p in reader.pages)
            if not text.strip():
                return "needs_transcription", 0
        else:
            if b"\0" in raw:
                return "invalid_type", 0
            text = raw.decode("utf-8-sig")
        if SECRET.search(text):
            return "sensitive_excluded", 0
        return "parsed", len(text)
    except Exception:
        return "parse_failed", 0


def scan(root: Path) -> dict:
    if not root.is_dir():
        raise BusinessError("INPUT_INVALID", "接入目录不存在或不是文件夹")
    files: list[FileRecord] = []
    warnings: list[str] = []
    total, entries = 0, 0
    try:
        for current, dirs, names, directory_fd in os.fwalk(root, follow_symlinks=False):
            dirs[:] = sorted(d for d in dirs if d not in EXCLUDED and not d.startswith("."))
            for d in dirs[:]:
                if (Path(current) / d).is_symlink():
                    dirs.remove(d)
                    warnings.append("已排除符号链接目录")
            entries += len(dirs) + len(names)
            if entries > MAX_ENTRIES:
                raise BusinessError("SCAN_LIMIT", "目录过大，请选择较小的资料目录")
            for name in sorted(names):
                path = Path(current) / name
                if name.startswith(".") or path.suffix.lower() in {".pem", ".key", ".p12"}:
                    continue
                if path.suffix.lower() not in FORMATS and name not in DEPENDENCIES:
                    continue
                if path.is_symlink() or not path.resolve().is_relative_to(root):
                    warnings.append("已排除符号链接或越界文件")
                    continue
                if len(files) >= MAX_FILES:
                    raise BusinessError("SCAN_LIMIT", "资料超过 500 项，请分批接入")
                try:
                    raw = read_safe(path, directory_fd)
                except BusinessError:
                    raise
                except OSError:
                    raise BusinessError("SCAN_CHANGED", "资料读取失败，请检查权限后重新预览") from None
                total += len(raw)
                if total > MAX_BATCH:
                    raise BusinessError("SCAN_LIMIT", "资料合计超过 100 MB，请分批接入")
                status, chars = parse(raw, path.suffix.lower())
                files.append(
                    {
                        "path": path.relative_to(root).as_posix(),
                        "size": len(raw),
                        "sha256": hashlib.sha256(raw).hexdigest(),
                        "status": status,
                        "chars": chars,
                    }
                )
    except OSError:
        raise BusinessError("SCAN_CHANGED", "无法读取资料目录，请检查权限") from None
    if not any(Path(f["path"]).name.lower().startswith("readme") for f in files):
        warnings.append("缺少 README")
    if not any("prd" in f["path"].lower() or "需求" in f["path"] for f in files):
        warnings.append("未识别到 PRD，后续需要补全需求")
    git_present = (root / ".git").exists()
    if git_present:
        warnings.append("发现 Git 元数据；本轮未执行仓库命令，未提交修改待后续核对")
    return {
        "files": files,
        "warnings": sorted(set(warnings)),
        "git_present": git_present,
        "git_changes": "not_checked",
    }
